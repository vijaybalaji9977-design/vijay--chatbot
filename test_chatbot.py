import sys
import io
import base64
from PIL import Image

# Force UTF-8 stdout if needed
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')

from fastapi.testclient import TestClient
from app import app
from core.intent import detect_intent, resolve_mode
from services.offline_engine import offline_engine
from services.document_service import document_service
from services.image_service import image_service

client = TestClient(app)

def test_health_and_modes():
    print("\n--- 1. Testing Health & Modes Endpoint ---")
    res = client.get("/api/health")
    assert res.status_code == 200, f"Health check failed: {res.text}"
    data = res.json()
    assert data["status"] == "healthy"
    assert "problem_solving" in data["features"]
    assert "image_understanding" in data["features"]
    print("[PASS] /api/health:", data["app"], "Version:", data["version"])

    res = client.get("/api/modes")
    assert res.status_code == 200, f"Modes endpoint failed: {res.text}"
    modes = res.json()["modes"]
    assert len(modes) >= 10, f"Expected at least 10 modes, got {len(modes)}"
    mode_ids = [m["id"] for m in modes]
    for required in ["general", "study", "coding", "math", "summary", "translate", "content", "vision", "document", "recommendation"]:
        assert required in mode_ids, f"Mode '{required}' missing from modes list!"
    print(f"[PASS] /api/modes verified: {len(modes)} modes available ({', '.join(mode_ids)}).")

def test_intent_detection():
    print("\n--- 2. Testing Intent Detection ---")
    test_cases = [
        ("solve x^2 - 5x + 6 = 0", "math"),
        ("differentiate x^3 + 2x with respect to x", "math"),
        ("write a python binary search algorithm", "coding"),
        ("translate hello world into Spanish", "translate"),
        ("summarize this article about AI in 3 bullet points", "summary"),
        ("explain quantum computing to a 10 year old", "study"),
        ("draft a professional cover letter for software engineer", "content"),
        ("recommend top 5 books on system design", "recommendation"),
        ("what is shown in this image?", "vision")
    ]
    for query, expected_mode in test_cases:
        detected_mode, conf = detect_intent(query)
        print(f"Query: '{query}' -> Detected: {detected_mode} (conf: {conf})")
        assert detected_mode == expected_mode, f"Expected {expected_mode}, got {detected_mode}"
    print("[PASS] All intent classifications passed!")

def test_offline_math_engine():
    print("\n--- 3. Testing SymPy Basic Problem Solving Solver ---")
    # Test derivative
    deriv_res = offline_engine.process_query("differentiate x^3 + 4*x^2 - 7*x + 12", mode="math")
    assert "3 x^{2} + 8 x - 7" in deriv_res or "3*x^2" in deriv_res or "3 x" in deriv_res, f"Unexpected derivative result: {deriv_res}"
    print("[PASS] Derivative calculus solution verified.")

    # Test equation
    eq_res = offline_engine.process_query("solve x^2 - 5*x + 6 = 0", mode="math")
    assert "2" in eq_res and "3" in eq_res, f"Unexpected equation result: {eq_res}"
    print("[PASS] Algebraic quadratic equation solver verified.")

    # Test arithmetic
    eval_res = offline_engine.process_query("sqrt(144) * 10", mode="math")
    assert "120" in eval_res, f"Unexpected calculation result: {eval_res}"
    print("[PASS] Symbolic calculation verified.")

def test_document_processing():
    print("\n--- 4. Testing Document Upload & Extraction ---")
    sample_text = "The AI-Powered Multi-Functional Chatbot is an intelligent conversational platform combining NLP, machine learning, and multi-mode assistance."
    doc_file = io.BytesIO(sample_text.encode("utf-8"))
    
    res = client.post("/api/upload", files={"file": ("test_project.txt", doc_file, "text/plain")})
    assert res.status_code == 200, f"Upload failed: {res.text}"
    doc_info = res.json()["document"]
    doc_id = doc_info["id"]
    print("[PASS] Uploaded document:", doc_info["filename"], "ID:", doc_id)

    # Test Q&A against document
    chat_res = client.post("/api/chat", json={
        "message": "What is the primary objective of this project?",
        "document_id": doc_id,
        "mode": "document",
        "provider": "offline"
    })
    assert chat_res.status_code == 200, f"Document chat failed: {chat_res.text}"
    print("[PASS] Document Q&A Response received successfully.")

def test_image_processing_and_vision():
    print("\n--- 5. Testing Image Upload & Multimodal Vision ---")
    # Generate a synthetic 100x100 RGB image
    img = Image.new("RGB", (100, 100), color=(73, 109, 137))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()

    # Upload via /api/upload
    res = client.post("/api/upload", files={"file": ("test_diagram.png", io.BytesIO(img_bytes), "image/png")})
    assert res.status_code == 200, f"Image upload failed: {res.text}"
    data = res.json()
    assert data["type"] == "image"
    img_data_url = data["image"]["data_url"]
    print("[PASS] Uploaded image via /api/upload:", data["image"]["filename"], f"({data['image']['width']}x{data['image']['height']})")

    # Chat with image attached
    chat_res = client.post("/api/chat", json={
        "message": "Analyze this diagram and describe its visual properties.",
        "image_data": img_data_url,
        "image_filename": "test_diagram.png",
        "mode": "vision",
        "provider": "offline"
    })
    assert chat_res.status_code == 200, f"Image chat failed: {chat_res.text}"
    reply = chat_res.json()["assistant_message"]["content"]
    assert "Image" in reply or "Visual" in reply or "Resolution" in reply
    print("[PASS] Image Understanding Q&A response received successfully.")

def test_chat_and_conversations():
    print("\n--- 6. Testing Chat, History & Feedback API ---")
    # Send a coding chat query
    chat_res = client.post("/api/chat", json={
        "message": "Write a python function for fibonacci sequence",
        "mode": "coding",
        "provider": "offline"
    })
    assert chat_res.status_code == 200, f"Chat failed: {chat_res.text}"
    data = chat_res.json()
    conv_id = data["conversation_id"]
    msg_id = data["assistant_message"]["id"]
    print(f"[PASS] Chat response generated for conversation {conv_id}.")
    print(f"       Mode resolved: {data['mode']}, Model: {data['model']}")

    # Submit feedback
    fb_res = client.post("/api/feedback", json={
        "conversation_id": conv_id,
        "message_id": msg_id,
        "rating": 1,
        "comment": "Great Python implementation!",
        "mode": "coding"
    })
    assert fb_res.status_code == 200, f"Feedback failed: {fb_res.text}"
    print("[PASS] Feedback recorded successfully.")

    # Get conversation details
    conv_get = client.get(f"/api/conversations/{conv_id}")
    assert conv_get.status_code == 200
    assert len(conv_get.json()["messages"]) >= 2
    print(f"[PASS] Conversation details retrieved with {len(conv_get.json()['messages'])} messages.")

def test_docx_processing():
    print("\n--- 7. Testing DOCX Document Upload & Parsing ---")
    import docx
    doc = docx.Document()
    doc.add_heading("Machine Learning Study Notes", 0)
    doc.add_paragraph("Supervised learning trains an algorithm on labeled data pairs.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    res = client.post("/api/upload", files={"file": ("study_notes.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, f"DOCX upload failed: {res.text}"
    doc_data = res.json()["document"]
    assert doc_data["file_type"] == "docx"
    assert "Supervised learning" in doc_data["preview"]
    print("[PASS] DOCX document uploaded & parsed cleanly:", doc_data["filename"])

def test_chat_streaming():
    print("\n--- 8. Testing SSE Streaming Chat API ---")
    res = client.post("/api/chat/stream", json={
        "message": "Explain quicksort in 2 sentences",
        "mode": "coding",
        "provider": "offline"
    })
    assert res.status_code == 200
    assert "text/event-stream" in res.headers.get("content-type", "")
    chunks = [line for line in res.iter_lines() if line]
    assert len(chunks) > 0, "No stream events received"
    print(f"[PASS] SSE streaming endpoint verified with {len(chunks)} stream lines.")

def test_static_index():
    print("\n--- 9. Testing Frontend Static Delivery ---")
    res = client.get("/")
    assert res.status_code == 200
    assert "AI-Powered Multi-Functional Chatbot" in res.text
    print("[PASS] Frontend HTML delivered cleanly.")

if __name__ == "__main__":
    print("========================================")
    print(" RUNNING AUTOMATED CHATBOT VERIFICATION ")
    print("========================================")
    test_health_and_modes()
    test_intent_detection()
    test_offline_math_engine()
    test_document_processing()
    test_image_processing_and_vision()
    test_chat_and_conversations()
    test_docx_processing()
    test_chat_streaming()
    test_static_index()
    print("\n========================================")
    print("  ALL 9 TESTS COMPLETED SUCCESSFULLY!   ")
    print("========================================")
